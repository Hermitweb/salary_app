"""
下拉框联动测试
- 验证 salary.html 的 year/month 下拉框正确显示
- 验证年份变化时月份下拉框正确更新
"""
import os
import sys
import re
import time
import tempfile
import shutil
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


@pytest.fixture(scope="module")
def test_env():
    """每个测试模块创建独立的临时数据目录和服务"""
    TMP_DATA = Path(tempfile.mkdtemp(prefix="dropdown_test_"))
    print(f"[setup] 临时数据目录: {TMP_DATA}")

    from control_panel_pkg import constants
    import excel_data
    import cache as cache_mod

    # 关键：先重置 DataCache 单例，再设置 DATA_DIR
    cache_mod.DataCache.reset_for_test()
    constants.DATA_DIR = TMP_DATA
    constants.DATA_DIR.mkdir(exist_ok=True)
    constants.PORT_CONFIG = TMP_DATA / "port.txt"
    constants.FIELDS_CONFIG = TMP_DATA / "fields.txt"
    constants.BASE_DIR = ROOT
    excel_data.DATA_DIR = str(TMP_DATA)
    excel_data.FIELDS_CONFIG = TMP_DATA / "fields.txt"

    # 准备员工表
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["员工号", "姓名", "部门", "紧急联系人号码"])
    ws.append(["E0001", "张三", "测试部", "13800000001"])
    wb.save(TMP_DATA / "员工列表.xlsx")

    # 准备多个月份的工资表（不同年）
    salary_dir = TMP_DATA / "salary"
    salary_dir.mkdir(exist_ok=True)
    salary_files = [
        ("2025年11月员工工资.xlsx", [("E0001", "张三", "测试部", 4000)]),
        ("2025年12月员工工资.xlsx", [("E0001", "张三", "测试部", 4500)]),
        ("2026年01月员工工资.xlsx", [("E0001", "张三", "测试部", 5000)]),
        ("2026年04月员工工资.xlsx", [("E0001", "张三", "测试部", 5500)]),
        ("2026年05月员工工资.xlsx", [("E0001", "张三", "测试部", 6000)]),
    ]
    for fname, rows in salary_files:
        w = Workbook()
        s = w.active
        s.append(["员工号", "姓名", "部门", "应发基本工资"])
        for r in rows:
            s.append(r)
        w.save(salary_dir / fname)
    # 设置端口
    PORT = 15020
    constants.PORT_CONFIG.write_text(str(PORT), encoding="utf-8")

    # 启动服务
    from control_panel_pkg.service import ServiceManager
    import logging
    logging.getLogger("werkzeug").setLevel(logging.WARNING)
    sm = ServiceManager()
    ok, msg = sm.start()
    assert ok, f"启动失败: {msg}"
    print(f"[start] {msg}")
    time.sleep(0.5)

    yield {
        "TMP_DATA": TMP_DATA,
        "salary_dir": salary_dir,
        "PORT": PORT,
        "sm": sm,
    }

    # 清理
    sm.stop()
    time.sleep(0.3)
    shutil.rmtree(TMP_DATA, ignore_errors=True)
    cache_mod.DataCache.reset_for_test()


def _login_and_get_salary_page(test_env, year=None, month=None):
    """登录并访问工资页"""
    import requests
    PORT = test_env["PORT"]
    BASE = f"http://127.0.0.1:{PORT}"
    s = requests.Session()
    r = s.get(f"{BASE}/login")
    m = re.search(r'name="_csrf_token"\s+value="([^"]+)"', r.text)
    csrf = m.group(1) if m else ""
    s.post(f"{BASE}/login",
           data={"username": "E0001", "password": "13800000001", "_csrf_token": csrf})
    params = {}
    if year is not None:
        params["year"] = year
    if month is not None:
        params["month"] = month
    r = s.get(f"{BASE}/salary", params=params)
    return r.text


def _extract_year_options(html):
    m = re.search(r'<select name="year"[^>]*>(.*?)</select>', html, re.DOTALL)
    if not m:
        return None
    opts = re.findall(r'<option\s+value="(\d+)"', m.group(1))
    return [int(x) for x in opts]


def _extract_month_options(html):
    m = re.search(r'<select name="month"[^>]*>(.*?)</select>', html, re.DOTALL)
    if not m:
        return None
    opts = re.findall(r'<option\s+value="(\d+)"', m.group(1))
    return [int(x) for x in opts]


def _extract_selected_year(html):
    m = re.search(r'<select name="year"[^>]*>.*?<option\s+value="(\d+)"[^>]*selected', html, re.DOTALL)
    return int(m.group(1)) if m else None


def _extract_selected_month(html):
    m = re.search(r'<select name="month"[^>]*>.*?<option\s+value="(\d+)"[^>]*selected', html, re.DOTALL)
    return int(m.group(1)) if m else None


def test_default(test_env):
    print("\n" + "="*60)
    print("T1: 默认访问（无参数）")
    print("="*60)
    html = _login_and_get_salary_page(test_env)
    years = _extract_year_options(html)
    months_opts = _extract_month_options(html)
    sel_year = _extract_selected_year(html)
    sel_month = _extract_selected_month(html)
    print(f"年份下拉框: {years}")
    print(f"月份下拉框: {months_opts}")
    print(f"默认选中: {sel_year}年 {sel_month}月")
    assert years == [2026, 2025], f"年份错误: {years}"
    assert sel_year == 2026
    assert sel_month == 5, f"应该选最新年最新月 5，实际 {sel_month}"
    assert months_opts == [5, 4, 1], f"2026年月份错误: {months_opts}"


def test_year_2025(test_env):
    print("\n" + "="*60)
    print("T2: 选 2025 年（应显示 11/12 两个月）")
    print("="*60)
    html = _login_and_get_salary_page(test_env, year=2025)
    years = _extract_year_options(html)
    months = _extract_month_options(html)
    sel_year = _extract_selected_year(html)
    sel_month = _extract_selected_month(html)
    print(f"年份下拉框: {years}")
    print(f"月份下拉框: {months}")
    print(f"默认选中: {sel_year}年 {sel_month}月")
    assert years == [2026, 2025]
    assert sel_year == 2025
    assert months == [12, 11], f"2025年月份错误: {months}"
    assert sel_month == 12, f"应该选 12月，实际 {sel_month}"


def test_specific_month(test_env):
    print("\n" + "="*60)
    print("T3: 选 2025年11月（应正常显示该月工资）")
    print("="*60)
    html = _login_and_get_salary_page(test_env, year=2025, month=11)
    sel_year = _extract_selected_year(html)
    sel_month = _extract_selected_month(html)
    print(f"选中: {sel_year}年 {sel_month}月")
    assert "应发基本工资" in html or "工资条" in html
    assert sel_year == 2025
    assert sel_month == 11


def test_invalid_month(test_env):
    print("\n" + "="*60)
    print("T4: 选了无效的 2026年13月（应回退到该年最新月 5）")
    print("="*60)
    html = _login_and_get_salary_page(test_env, year=2026, month=13)
    sel_month = _extract_selected_month(html)
    print(f"回退后选中月份: {sel_month}")
    assert sel_month == 5, f"应回退到 5，实际 {sel_month}"


def test_invalid_year(test_env):
    print("\n" + "="*60)
    print("T5: 选了无效的 2030年5月（应回退到最新年 2026）")
    print("="*60)
    html = _login_and_get_salary_page(test_env, year=2030, month=5)
    sel_year = _extract_selected_year(html)
    sel_month = _extract_selected_month(html)
    print(f"回退后选中: {sel_year}年 {sel_month}月")
    assert sel_year == 2026, f"应回退到 2026，实际 {sel_year}"
    assert sel_month == 5


def test_concurrent(test_env):
    port = test_env["PORT"]
    print("\n" + "="*60)
    print("T6: 20 并发访问下拉框（缓存）")
    print("="*60)
    start = time.time()
    with ThreadPoolExecutor(max_workers=20) as pool:
        results = [f.result() for f in [
            pool.submit(_login_and_get_salary_page, test_env, 2026, 5) for _ in range(20)
        ]]
    elapsed = time.time() - start
    for i, html in enumerate(results):
        sel_year = _extract_selected_year(html)
        sel_month = _extract_selected_month(html)
        assert sel_year == 2026 and sel_month == 5, f"第{i}个请求异常: {sel_year}-{sel_month}"
    print(f"20 并发完成, 耗时 {elapsed*1000:.0f}ms")


def test_mtime_fallback():
    """测试文件名没有年月时，mtime 兜底逻辑"""
    import shutil
    import excel_data

    # 临时建一个独立数据目录
    tmp = Path(tempfile.mkdtemp(prefix="mtime_test_"))
    salary_dir = tmp / "salary"
    salary_dir.mkdir()

    from openpyxl import Workbook
    w = Workbook()
    s = w.active
    s.append(["员工号", "姓名", "部门", "应发基本工资"])
    s.append(["E9999", "测试员", "测试部", 5000])
    # 文件名没有年月
    fpath = salary_dir / "工资单.xlsx"
    w.save(fpath)
    # 设置 mtime 为 2026-03-15
    import datetime
    target_time = datetime.datetime(2026, 3, 15, 10, 0, 0).timestamp()
    os.utime(fpath, (target_time, target_time))

    # 修改 excel_data.DATA_DIR
    old_dir = excel_data.DATA_DIR
    excel_data.DATA_DIR = str(tmp)
    try:
        months = excel_data.get_available_months()
        print(f"\n[T7] mtime 兜底测试 - 文件名无年月")
        print(f"  get_available_months = {months}")
        assert len(months) == 1, f"应返回 1 条，实际 {len(months)}"
        assert months[0]["year"] == 2026, f"年份应为 2026，实际 {months[0]['year']}"
        assert months[0]["month"] == 3, f"月份应为 3，实际 {months[0]['month']}"
        print("  T7 PASS")
    finally:
        excel_data.DATA_DIR = old_dir
        shutil.rmtree(tmp, ignore_errors=True)


def test_employee_not_salary():
    """测试员工列表文件不会被误识别为工资表"""
    import shutil
    import excel_data

    tmp = Path(tempfile.mkdtemp(prefix="emp_not_sal_"))
    # 在 data 根目录（不是 salary/）放员工列表
    from openpyxl import Workbook
    w = Workbook()
    s = w.active
    s.append(["员工号", "姓名", "部门"])
    s.append(["E0001", "张三", "测试部"])
    w.save(tmp / "员工列表.xlsx")

    old_dir = excel_data.DATA_DIR
    excel_data.DATA_DIR = str(tmp)
    try:
        months = excel_data.get_available_months()
        print(f"\n[T8] 员工列表不应被识别为工资表")
        print(f"  get_available_months = {months}")
        assert len(months) == 0, f"员工列表不应被识别为工资表，实际: {months}"
        print("  T8 PASS")
    finally:
        excel_data.DATA_DIR = old_dir
        shutil.rmtree(tmp, ignore_errors=True)
