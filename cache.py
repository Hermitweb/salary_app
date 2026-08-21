"""
数据访问层缓存
- 员工表、工资表按 mtime 缓存到内存
- 线程安全（RLock）
- 懒加载：第一次访问时才加载
- 内存占用：30-100 个员工约 1-2MB

设计目标：
- 单次查询耗时从 650ms → < 5ms（纯内存查找）
- 30 并发用户 QPS 从 1.5 → 100+
"""
import os
import threading
from typing import Any, Optional, Tuple, List, Dict

try:
    import openpyxl
except ImportError:
    openpyxl = None

# 工资表缓存文件数上限（约 2 年的月份数）：超出按最久未使用淘汰，
# 防止长期运行中缓存随访问过的月份数无限增长
_MAX_SALARY_CACHE_FILES = 24

# 工号列表头关键词（子串匹配，兼容「员工工号」「账号」等变体）
_EMP_ID_KEYWORDS = ("工号", "员工号", "账号")


def find_salary_emp_id_col(headers):
    """在工资表表头中定位工号列（缓存层与 excel_data 共用，保证判定一致）。

    先精确匹配「员工号 / 工号」，再子串匹配变体（含「账号」）。
    找不到返回 None —— 绝不回落到第 0 列，防止误用其他列当工号建索引。
    """
    for i, h in enumerate(headers):
        if h in ("员工号", "工号"):
            return i
    for i, h in enumerate(headers):
        if h and any(kw in h for kw in _EMP_ID_KEYWORDS):
            return i
    return None


class DataCache:
    """全局单例缓存：员工表 + 各月工资表"""

    _instance: Optional["DataCache"] = None
    _init_lock = threading.Lock()

    def __init__(self):
        # 员工表
        self._emp_path: Optional[str] = None
        self._emp_mtime: Optional[float] = None
        self._emp_rows: List[Dict[str, Any]] = []
        self._emp_headers: List[str] = []
        # 列索引为 None 表示员工表中未找到对应列（绝不回落到第 0 列，防止误用工号列当密码列）
        self._emp_id_idx: Optional[int] = None
        self._emp_pwd_idx: Optional[int] = None
        self._emp_name_idx: Optional[int] = None
        self._emp_dept_idx: Optional[int] = None
        self._emp_status_idx: int = -1  # 离职列
        self._emp_locked = threading.RLock()

        # 工资表：path -> (mtime, headers, rows, [emp_id -> row_dict])
        self._salary_cache: Dict[str, Tuple[float, List[str], List[Dict[str, Any]], Dict[str, Dict[str, Any]]]] = {}
        self._salary_lock = threading.RLock()

        # 统计
        self._hits = 0
        self._misses = 0

    @classmethod
    def instance(cls) -> "DataCache":
        if cls._instance is None:
            with cls._init_lock:
                if cls._instance is None:
                    cls._instance = cls()
        return cls._instance

    @classmethod
    def reset_for_test(cls):
        """测试用：重置单例"""
        with cls._init_lock:
            cls._instance = None

    # ==================== 员工表 ====================
    def get_employees(self, path: str) -> Tuple[List[Dict[str, Any]], List[str], Optional[int], Optional[int], Optional[int], Optional[int], int]:
        """获取员工表数据（带 mtime 缓存）

        返回：(rows列表, headers, id_idx, pwd_idx, name_idx, dept_idx, status_idx)
        列索引可能为 None（表中未找到该列），调用方必须显式处理
        """
        with self._emp_locked:
            try:
                mtime = os.path.getmtime(path)
            except OSError:
                return [], [], None, None, None, None, -1

            if self._emp_path == path and self._emp_mtime == mtime and self._emp_rows:
                self._hits += 1
                return (
                    self._emp_rows, self._emp_headers,
                    self._emp_id_idx, self._emp_pwd_idx,
                    self._emp_name_idx, self._emp_dept_idx,
                    self._emp_status_idx,
                )

            # 缓存未命中或文件变更，重新加载
            self._misses += 1
            rows, headers, id_idx, pwd_idx, name_idx, dept_idx, status_idx = \
                self._load_employees_from_disk(path)
            self._emp_path = path
            self._emp_mtime = mtime
            self._emp_rows = rows
            self._emp_headers = headers
            self._emp_id_idx = id_idx
            self._emp_pwd_idx = pwd_idx
            self._emp_name_idx = name_idx
            self._emp_dept_idx = dept_idx
            self._emp_status_idx = status_idx
            return rows, headers, id_idx, pwd_idx, name_idx, dept_idx, status_idx

    @staticmethod
    def _load_employees_from_disk(path: str):
        """从磁盘加载员工表（耗时操作，约 50-200ms）"""
        if openpyxl is None:
            return [], [], None, None, None, None, -1
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                headers = list(next(rows_iter))
            except StopIteration:
                return [], [], None, None, None, None, -1

            # 检测列索引（找不到返回 None，由调用方决定如何处理）
            def find_idx(keywords):
                for kw in keywords:
                    for i, h in enumerate(headers):
                        if h is not None and kw in str(h):
                            return i
                return None

            id_idx = find_idx(["工号", "员工号", "账号"])
            pwd_idx = find_idx(["密码", "手机号", "联系电话",
                                "紧急联系人", "紧急联系人号码"])
            name_idx = find_idx(["姓名", "名字"])
            dept_idx = find_idx(["部门", "科室"])
            status_idx = find_idx(["离职", "在职状态", "状态"])

            # 转为 list[dict]
            rows = []
            for row in rows_iter:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                d = {}
                for i, h in enumerate(headers):
                    if h is None:
                        continue
                    if i < len(row):
                        d[str(h)] = row[i]
                rows.append(d)

            return rows, headers, id_idx, pwd_idx, name_idx, dept_idx, status_idx
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def invalidate_employees(self):
        """主动失效员工缓存（管理员更新员工表后调用）"""
        with self._emp_locked:
            self._emp_path = None
            self._emp_mtime = None
            self._emp_rows = []

    # ==================== 工资表 ====================
    def get_salary(self, path: str) -> Tuple[List[str], List[Dict[str, Any]], Dict[str, Dict[str, Any]]]:
        """获取工资表数据（带 mtime 缓存）

        返回：(headers, rows列表, emp_id->row 索引)
        """
        with self._salary_lock:
            try:
                mtime = os.path.getmtime(path)
            except OSError:
                return [], [], {}

            entry = self._salary_cache.get(path)
            if entry is not None and entry[0] == mtime:
                self._hits += 1
                # 标记为最近使用（移到字典末尾，dict 按插入序迭代即 LRU）
                if len(self._salary_cache) > 1:
                    self._salary_cache[path] = self._salary_cache.pop(path)
                return entry[1], entry[2], entry[3]

            # 重新加载
            self._misses += 1
            headers, rows, idx_map = self._load_salary_from_disk(path)
            self._salary_cache[path] = (mtime, headers, rows, idx_map)
            # 超出上限时淘汰最久未使用的条目（插入序最前者）
            while len(self._salary_cache) > _MAX_SALARY_CACHE_FILES:
                oldest = next(iter(self._salary_cache))
                self._salary_cache.pop(oldest, None)
            return headers, rows, idx_map

    @staticmethod
    def _load_salary_from_disk(path: str):
        """从磁盘加载工资表（耗时操作，约 300-600ms）"""
        if openpyxl is None:
            return [], [], {}
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                headers = list(next(rows_iter))
            except StopIteration:
                return [], [], {}

            headers = [str(h) if h is not None else "" for h in headers]

            # 找工号列（与 excel_data.get_salary 共用同一判定；
            # 找不到时不建索引，绝不回落到第 0 列）
            emp_id_idx = find_salary_emp_id_col(headers)

            rows = []
            idx_map = {}
            for row in rows_iter:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                d = {}
                for i, h in enumerate(headers):
                    if i < len(row):
                        d[h] = row[i]
                rows.append(d)
                if emp_id_idx is not None:
                    emp_id = d.get(headers[emp_id_idx])
                    if emp_id is not None:
                        idx_map[str(emp_id).strip()] = d

            return headers, rows, idx_map
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def invalidate_salary(self, path: Optional[str] = None):
        """主动失效工资缓存；不传 path 则清空全部"""
        with self._salary_lock:
            if path is None:
                self._salary_cache.clear()
            else:
                self._salary_cache.pop(path, None)

    def invalidate_all(self):
        """主动失效所有缓存"""
        self.invalidate_employees()
        self.invalidate_salary(None)

    def get_stats(self) -> Dict[str, int]:
        """获取缓存统计"""
        return {
            "hits": self._hits,
            "misses": self._misses,
            "salary_files_cached": len(self._salary_cache),
        }
